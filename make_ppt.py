# make_ppt.py
# Builds the college presentation for the Gold Price Prediction project.
# Produces BOTH a real PowerPoint (.pptx) and a PDF mirror (for older Office PCs).
#   python make_ppt.py
import os

HERE = os.path.dirname(os.path.abspath(__file__))
DOCS = os.path.join(HERE, "docs")
LOGO = os.path.join(HERE, "assets", "college_logo.jpg")
LOGO_NAVY = os.path.join(HERE, "assets", "college_logo_navy.png")
SHOT_PREDICT = os.path.join(DOCS, "screenshot_predict.png")
SHOT_TRAIN = os.path.join(DOCS, "screenshot_train.png")
RESULTS_CHART = os.path.join(DOCS, "results_chart.png")

TITLE = "Gold Price Prediction using Machine Learning"
STUDENT = "Sanchay Bhardwaj"
ROLL = "2300650100104"
GUIDE = "Mrs. Ashna Thakur"
COLLEGE = "B.S.A. College of Engineering & Technology, Mathura"
DEPT = "Dept. of Computer Science & Engineering   |   Session 2026-27"

# model, MAE, RMSE, R2  (matches model/metrics.json)
MODELS = [
    ("RandomForest", 32.71, 44.47, 0.9613),
    ("Linear Regression", 32.91, 44.60, 0.9611),
    ("Ridge", 33.33, 45.50, 0.9595),
    ("Naive baseline", 36.14, 47.97, 0.9549),
]
BEST = "RandomForest"

SLIDES = [
    {"kind": "title"},
    {"kind": "bullets", "title": "Introduction", "bullets": [
        "Gold is a globally traded precious metal and a traditional safe-haven asset.",
        "Its price is affected by demand, inflation, currency value and market sentiment.",
        "Being able to forecast gold prices helps investors and analysts take better decisions.",
        "This project builds a Machine Learning model that predicts the monthly price of gold.",
    ]},
    {"kind": "bullets", "title": "Problem Statement", "bullets": [
        "Given the historical monthly price of gold, predict its price for the coming months.",
        "This is a supervised regression problem: past behaviour is the input, future price is the output.",
        "Multiple models are trained and the most accurate one is kept for prediction.",
        "The winning model is then wrapped in a simple desktop application.",
    ]},
    {"kind": "bullets", "title": "Objectives", "bullets": [
        "Prepare a clean monthly gold price dataset for the years 2000 to 2024.",
        "Engineer features that capture the trend and the short-term momentum of prices.",
        "Train and compare RandomForest, Linear Regression and Ridge models.",
        "Measure accuracy using MAE, RMSE and R-squared against a naive baseline.",
        "Provide a simple desktop GUI so anyone can run a prediction.",
    ]},
    {"kind": "bullets", "title": "Dataset", "bullets": [
        "Monthly gold price in USD per troy ounce, from 2000 to 2024 (300 records).",
        "Columns in the file: Date, Year, Month and Price.",
        "Built from public annual-average anchors with a small momentum effect added.",
        "Model features: the last 3 monthly returns, plus the month and the year.",
    ]},
    {"kind": "bullets", "title": "Methodology", "bullets": [
        "The model predicts the month-to-month RETURN, not the raw price.",
        "Why: tree models cannot guess a price higher than they have seen, so a rising trend breaks them.",
        "Price is rebuilt as   next price = last price x (1 + predicted return).",
        "Forecasting rolls forward one month at a time, feeding each answer back as the next input.",
        "A chronological split is used - the last 60 months are kept for testing.",
    ]},
    {"kind": "bullets", "title": "Technology Used", "bullets": [
        "Language: Python.",
        "Data handling: pandas and NumPy.",
        "Models: scikit-learn - RandomForest, Linear Regression, Ridge.",
        "Charts: Matplotlib.",
        "Interface: Tkinter.   Save / Load: Joblib.",
    ]},
    {"kind": "table", "title": "Results"},
    {"kind": "image", "title": "Application - Prediction", "image": SHOT_PREDICT,
     "caption": "Desktop GUI: choose a horizon and forecast the gold price month by month."},
    {"kind": "image", "title": "Application - Model Comparison", "image": SHOT_TRAIN,
     "caption": "Train tab: accuracy table and the error comparison across all models."},
    {"kind": "bullets", "title": "Conclusion", "bullets": [
        "RandomForest gave the best accuracy and clearly beat the naive baseline.",
        "Predicting returns instead of raw prices made the models reliable on a trending series.",
        "The Tkinter GUI turns the trained model into an easy-to-use forecasting tool.",
        "The project demonstrates a complete Machine Learning pipeline from data to interface.",
    ]},
    {"kind": "bullets", "title": "Future Scope", "bullets": [
        "Add real economic features such as inflation, USD index and interest rates.",
        "Try deep-learning time-series models like LSTM for longer forecasts.",
        "Connect a live gold price API so predictions always use current data.",
        "Show a confidence range instead of a single predicted value.",
    ]},
    {"kind": "end"},
]


def make_results_chart():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    names = [m[0] for m in MODELS]
    maes = [m[1] for m in MODELS]
    colors = ["#C9A227" if n == BEST else "#1F3864" for n in names]
    fig, ax = plt.subplots(figsize=(5.2, 4.0), dpi=130)
    bars = ax.bar(names, maes, color=colors)
    ax.set_title("Mean Absolute Error by model", color="#1F3864", fontsize=13, fontweight="bold")
    ax.set_ylabel("MAE (USD)")
    ax.tick_params(axis="x", labelsize=8, rotation=12)
    for b, v in zip(bars, maes):
        ax.text(b.get_x() + b.get_width() / 2, v, "%.1f" % v, ha="center", va="bottom", fontsize=9)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(RESULTS_CHART, facecolor="white")
    plt.close(fig)


# ------------------------------------------------------------------ PPTX
def render_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE

    NAVY = RGBColor(0x1F, 0x38, 0x64)
    GOLD = RGBColor(0xC9, 0xA2, 0x27)
    DARK = RGBColor(0x22, 0x22, 0x22)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    MUT = RGBColor(0x6B, 0x72, 0x80)
    LIGHT = RGBColor(0xEE, 0xF3, 0xFA)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    def rect(slide, l, t, w, h, color):
        s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
        s.fill.solid(); s.fill.fore_color.rgb = color
        s.line.fill.background(); s.shadow.inherit = False
        return s

    def tb(slide, l, t, w, h, anchor=None):
        box = slide.shapes.add_textbox(l, t, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        if anchor is not None:
            tf.vertical_anchor = anchor
        return tf

    def para(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, first=False, after=6):
        p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
        r.font.name = "Calibri"
        return p

    def header(slide, title):
        rect(slide, 0, 0, W, Inches(1.05), NAVY)
        rect(slide, 0, Inches(1.05), W, Inches(0.07), GOLD)
        tf = tb(slide, Inches(0.5), Inches(0.15), W - Inches(4.2), Inches(0.8), MSO_ANCHOR.MIDDLE)
        para(tf, title, 26, WHITE, bold=True, first=True)
        tf2 = tb(slide, W - Inches(4.0), Inches(0.15), Inches(3.5), Inches(0.8), MSO_ANCHOR.MIDDLE)
        para(tf2, "Gold Price Prediction", 12, GOLD, bold=True, align=PP_ALIGN.RIGHT, first=True)

    def footer(slide, n):
        rect(slide, 0, H - Inches(0.42), W, Inches(0.42), NAVY)
        rect(slide, 0, H - Inches(0.42), Inches(0.14), Inches(0.42), GOLD)
        tf = tb(slide, Inches(0.3), H - Inches(0.42), W - Inches(2.0), Inches(0.42), MSO_ANCHOR.MIDDLE)
        para(tf, "%s (%s)  |  %s" % (STUDENT, ROLL, COLLEGE), 9, WHITE, first=True)
        tf2 = tb(slide, W - Inches(1.6), H - Inches(0.42), Inches(1.3), Inches(0.42), MSO_ANCHOR.MIDDLE)
        para(tf2, "Slide %d" % n, 9, GOLD, bold=True, align=PP_ALIGN.RIGHT, first=True)

    for idx, sl in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        kind = sl["kind"]

        if kind == "title":
            rect(slide, 0, 0, W, H, NAVY)
            rect(slide, 0, Inches(4.35), W, Inches(0.06), GOLD)
            slide.shapes.add_picture(LOGO_NAVY, int((W - Inches(1.9)) / 2), Inches(0.6), Inches(1.9), Inches(1.9))
            tf = tb(slide, Inches(0.8), Inches(2.55), W - Inches(1.6), Inches(1.6))
            para(tf, TITLE, 34, WHITE, bold=True, align=PP_ALIGN.CENTER, first=True, after=6)
            para(tf, "A Mini Project in Machine Learning", 18, GOLD, align=PP_ALIGN.CENTER)
            tf2 = tb(slide, Inches(1.0), Inches(4.7), W - Inches(2.0), Inches(2.4))
            para(tf2, "Submitted by", 13, GOLD, bold=True, align=PP_ALIGN.CENTER, first=True, after=2)
            para(tf2, "%s  (%s)" % (STUDENT, ROLL), 22, WHITE, bold=True, align=PP_ALIGN.CENTER, after=14)
            para(tf2, "Project Guide:  %s" % GUIDE, 15, WHITE, align=PP_ALIGN.CENTER, after=4)
            para(tf2, COLLEGE, 14, WHITE, align=PP_ALIGN.CENTER, after=4)
            para(tf2, DEPT, 12, GOLD, align=PP_ALIGN.CENTER)

        elif kind == "bullets":
            header(slide, sl["title"])
            footer(slide, idx)
            rect(slide, Inches(0.55), Inches(1.7), Inches(0.09), Inches(4.7), GOLD)
            tf = tb(slide, Inches(0.95), Inches(1.5), W - Inches(1.6), H - Inches(2.4), MSO_ANCHOR.MIDDLE)
            first = True
            for b in sl["bullets"]:
                para(tf, "•   " + b, 20, DARK, first=first, after=22)
                first = False

        elif kind == "table":
            header(slide, sl["title"])
            footer(slide, idx)
            rows = len(MODELS) + 1
            gt = slide.shapes.add_table(rows, 4, Inches(0.7), Inches(1.9),
                                        Inches(7.3), Inches(3.4)).table
            hdr = ["Model", "MAE (USD)", "RMSE (USD)", "R-squared"]
            for c, h in enumerate(hdr):
                cell = gt.cell(0, c)
                cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
                cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                para(cell.text_frame, h, 15, WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
            for r, m in enumerate(MODELS, start=1):
                vals = [m[0], "%.2f" % m[1], "%.2f" % m[2], "%.4f" % m[3]]
                is_best = (m[0] == BEST)
                for c, val in enumerate(vals):
                    cell = gt.cell(r, c)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = LIGHT if r % 2 else WHITE
                    cell.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                    col = GOLD if is_best else DARK
                    para(cell.text_frame, val, 14, col, bold=is_best,
                         align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER, first=True)
            slide.shapes.add_picture(RESULTS_CHART, Inches(8.4), Inches(1.7), Inches(4.3), Inches(3.6))
            tf = tb(slide, Inches(0.7), Inches(5.7), Inches(7.3), Inches(1.0))
            para(tf, "RandomForest wins on every metric. Test window: last 60 months (from Jan 2020). "
                     "Lower MAE / RMSE and higher R-squared are better.", 12, MUT, first=True)

        elif kind == "image":
            header(slide, sl["title"])
            footer(slide, idx)
            from PIL import Image
            iw, ih = Image.open(sl["image"]).size
            max_w = W - Inches(1.6)
            max_h = H - Inches(2.6)
            scale = min(max_w / iw, max_h / ih)
            pw, ph = int(iw * scale), int(ih * scale)
            slide.shapes.add_picture(sl["image"], int((W - pw) / 2), Inches(1.45), pw, ph)
            tf = tb(slide, Inches(1.0), H - Inches(1.35), W - Inches(2.0), Inches(0.7))
            para(tf, sl["caption"], 13, MUT, align=PP_ALIGN.CENTER, first=True)

        elif kind == "end":
            rect(slide, 0, 0, W, H, NAVY)
            rect(slide, 0, Inches(3.55), W, Inches(0.06), GOLD)
            slide.shapes.add_picture(LOGO_NAVY, int((W - Inches(1.4)) / 2), Inches(0.9), Inches(1.4), Inches(1.4))
            tf = tb(slide, Inches(1.0), Inches(2.7), W - Inches(2.0), Inches(1.2))
            para(tf, "Thank You", 54, WHITE, bold=True, align=PP_ALIGN.CENTER, first=True)
            tf2 = tb(slide, Inches(1.0), Inches(4.0), W - Inches(2.0), Inches(2.6))
            para(tf2, "Questions & suggestions are welcome", 18, GOLD, align=PP_ALIGN.CENTER, first=True, after=16)
            para(tf2, "%s  (%s)" % (STUDENT, ROLL), 15, WHITE, align=PP_ALIGN.CENTER, after=3)
            para(tf2, "Guide: %s" % GUIDE, 14, WHITE, align=PP_ALIGN.CENTER, after=3)
            para(tf2, COLLEGE, 13, GOLD, align=PP_ALIGN.CENTER)

    prs.save(path)
    print("pptx:", path)


# ------------------------------------------------------------------ PDF
def render_pdf(path):
    from reportlab.lib.units import inch
    from reportlab.lib.colors import HexColor
    from reportlab.pdfgen import canvas as pdfcanvas
    from reportlab.platypus import Table, TableStyle
    from reportlab.lib import colors as rcolors

    NAVY = HexColor("#1F3864")
    GOLD = HexColor("#C9A227")
    DARK = HexColor("#222222")
    MUT = HexColor("#6B7280")
    WHITE = HexColor("#FFFFFF")
    LIGHT = HexColor("#EEF3FA")

    PW, PH = 13.333 * inch, 7.5 * inch
    c = pdfcanvas.Canvas(path, pagesize=(PW, PH))
    c.setTitle(TITLE)
    c.setAuthor(STUDENT)

    def wrap(text, font, size, maxw):
        words = text.split()
        lines, cur = [], ""
        for w in words:
            t = (cur + " " + w).strip()
            if c.stringWidth(t, font, size) <= maxw:
                cur = t
            else:
                lines.append(cur); cur = w
        if cur:
            lines.append(cur)
        return lines

    def header(title):
        c.setFillColor(NAVY); c.rect(0, PH - 1.05 * inch, PW, 1.05 * inch, fill=1, stroke=0)
        c.setFillColor(GOLD); c.rect(0, PH - 1.12 * inch, PW, 0.07 * inch, fill=1, stroke=0)
        c.setFillColor(WHITE); c.setFont("Helvetica-Bold", 24)
        c.drawString(0.5 * inch, PH - 0.68 * inch, title)
        c.setFillColor(GOLD); c.setFont("Helvetica-Bold", 12)
        c.drawRightString(PW - 0.5 * inch, PH - 0.66 * inch, "Gold Price Prediction")

    def footer(n):
        c.setFillColor(NAVY); c.rect(0, 0, PW, 0.42 * inch, fill=1, stroke=0)
        c.setFillColor(GOLD); c.rect(0, 0, 0.14 * inch, 0.42 * inch, fill=1, stroke=0)
        c.setFillColor(WHITE); c.setFont("Helvetica", 9)
        c.drawString(0.3 * inch, 0.16 * inch, "%s (%s)  |  %s" % (STUDENT, ROLL, COLLEGE))
        c.setFillColor(GOLD); c.setFont("Helvetica-Bold", 9)
        c.drawRightString(PW - 0.3 * inch, 0.16 * inch, "Slide %d" % n)

    def bullets(items):
        size, lead, gap = 19, 26, 16
        all_lines = []
        for b in items:
            all_lines.append(("•  " + b))
        # measure
        block = []
        total = 0
        for b in items:
            lines = wrap("•  " + b, "Helvetica", size, PW - 1.6 * inch)
            block.append(lines)
            total += len(lines) * lead + gap
        total -= gap
        top = PH - 1.15 * inch
        bottom = 0.55 * inch
        y = top - ((top - bottom) - total) / 2.0
        c.setFillColor(GOLD); c.rect(0.55 * inch, bottom, 0.09 * inch, top - bottom, fill=1, stroke=0)
        for lines in block:
            c.setFillColor(DARK); c.setFont("Helvetica", size)
            for ln in lines:
                c.drawString(0.95 * inch, y - size, ln)
                y -= lead
            y -= gap

    # ---- title slide ----
    c.setFillColor(NAVY); c.rect(0, 0, PW, PH, fill=1, stroke=0)
    c.drawImage(LOGO_NAVY, (PW - 1.9 * inch) / 2, PH - 0.6 * inch - 1.9 * inch, 1.9 * inch, 1.9 * inch)
    c.setFillColor(WHITE); c.setFont("Helvetica-Bold", 32)
    c.drawCentredString(PW / 2, PH - 3.35 * inch, TITLE)
    c.setFillColor(GOLD); c.setFont("Helvetica-Bold", 17)
    c.drawCentredString(PW / 2, PH - 3.85 * inch, "A Mini Project in Machine Learning")
    c.setFillColor(GOLD); c.rect(0, PH - 4.2 * inch, PW, 0.05 * inch, fill=1, stroke=0)
    c.setFillColor(GOLD); c.setFont("Helvetica-Bold", 13)
    c.drawCentredString(PW / 2, PH - 4.7 * inch, "Submitted by")
    c.setFillColor(WHITE); c.setFont("Helvetica-Bold", 22)
    c.drawCentredString(PW / 2, PH - 5.15 * inch, "%s  (%s)" % (STUDENT, ROLL))
    c.setFont("Helvetica", 14)
    c.drawCentredString(PW / 2, PH - 5.7 * inch, "Project Guide:  %s" % GUIDE)
    c.drawCentredString(PW / 2, PH - 6.05 * inch, COLLEGE)
    c.setFillColor(GOLD); c.setFont("Helvetica", 12)
    c.drawCentredString(PW / 2, PH - 6.4 * inch, DEPT)
    c.showPage()

    idx = 2
    for sl in SLIDES:
        kind = sl["kind"]
        if kind == "title":
            continue
        if kind == "bullets":
            header(sl["title"]); footer(idx); bullets(sl["bullets"])
            c.showPage(); idx += 1
        elif kind == "table":
            header(sl["title"]); footer(idx)
            data = [["Model", "MAE (USD)", "RMSE (USD)", "R-squared"]]
            for m in MODELS:
                data.append([m[0], "%.2f" % m[1], "%.2f" % m[2], "%.4f" % m[3]])
            t = Table(data, colWidths=[2.4 * inch, 1.4 * inch, 1.5 * inch, 1.5 * inch],
                      rowHeights=[0.6 * inch] + [0.55 * inch] * len(MODELS))
            style = [
                ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                ("TEXTCOLOR", (0, 0), (-1, 0), WHITE),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 12),
                ("FONTNAME", (0, 1), (0, 1), "Helvetica-Bold"),
                ("TEXTCOLOR", (0, 1), (0, 1), GOLD),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LIGHT]),
                ("GRID", (0, 0), (-1, -1), 0.6, HexColor("#C9D6EA")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ]
            t.setStyle(TableStyle(style))
            tw, th = t.wrap(7.3 * inch, 4 * inch)
            t.drawOn(c, 0.7 * inch, PH - 1.4 * inch - th)
            c.drawImage(RESULTS_CHART, 8.5 * inch, PH - 1.6 * inch - 3.6 * inch,
                        4.2 * inch, 3.6 * inch)
            note = wrap("RandomForest wins on every metric. Test window: last 60 months (from Jan 2020). "
                        "Lower MAE / RMSE and higher R-squared are better.", "Helvetica-Oblique", 11, 7.3 * inch)
            c.setFillColor(MUT); c.setFont("Helvetica-Oblique", 11)
            ny = PH - 1.4 * inch - th - 0.5 * inch
            for ln in note:
                c.drawString(0.7 * inch, ny, ln); ny -= 15
            c.showPage(); idx += 1
        elif kind == "image":
            header(sl["title"]); footer(idx)
            from PIL import Image as PImage
            iw, ih = PImage.open(sl["image"]).size
            max_w, max_h = PW - 1.6 * inch, PH - 2.7 * inch
            scale = min(max_w / iw, max_h / ih)
            pw, ph = iw * scale, ih * scale
            c.drawImage(sl["image"], (PW - pw) / 2, PH - 1.35 * inch - ph, pw, ph)
            c.setFillColor(MUT); c.setFont("Helvetica", 12)
            c.drawCentredString(PW / 2, 0.7 * inch, sl["caption"])
            c.showPage(); idx += 1
        elif kind == "end":
            c.setFillColor(NAVY); c.rect(0, 0, PW, PH, fill=1, stroke=0)
            c.drawImage(LOGO_NAVY, (PW - 1.4 * inch) / 2, PH - 0.9 * inch - 1.4 * inch, 1.4 * inch, 1.4 * inch)
            c.setFillColor(WHITE); c.setFont("Helvetica-Bold", 50)
            c.drawCentredString(PW / 2, PH - 3.1 * inch, "Thank You")
            c.setFillColor(GOLD); c.rect(0, PH - 3.4 * inch, PW, 0.05 * inch, fill=1, stroke=0)
            c.setFont("Helvetica-Bold", 17)
            c.drawCentredString(PW / 2, PH - 4.1 * inch, "Questions & suggestions are welcome")
            c.setFillColor(WHITE); c.setFont("Helvetica", 15)
            c.drawCentredString(PW / 2, PH - 4.8 * inch, "%s  (%s)" % (STUDENT, ROLL))
            c.setFont("Helvetica", 14)
            c.drawCentredString(PW / 2, PH - 5.15 * inch, "Guide: %s" % GUIDE)
            c.setFillColor(GOLD); c.setFont("Helvetica", 13)
            c.drawCentredString(PW / 2, PH - 5.5 * inch, COLLEGE)
            c.showPage(); idx += 1

    c.save()
    print("pdf:", path)


if __name__ == "__main__":
    os.makedirs(DOCS, exist_ok=True)
    make_results_chart()
    render_pptx(os.path.join(DOCS, "Gold-Price-Prediction-Presentation.pptx"))
    render_pdf(os.path.join(DOCS, "Gold-Price-Prediction-Presentation.pdf"))
